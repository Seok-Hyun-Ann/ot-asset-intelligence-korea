# -*- coding: utf-8 -*-
"""보고서를 엑셀로 (ADR-049).

**바이트 동일하게 쓴다.** 같은 입력이면 같은 파일이 나오고, 감사인이 다시 돌려
해시를 맞춰 볼 수 있다. 그래서 두 가지를 고정한다:

1. **zip 멤버 수정 시각** — `openpyxl` 은 '지금' 을 박는다. 1980-01-01 로 고정한다
   (재현 가능한 빌드의 표준 선택).
2. **`docProps/core.xml` 의 `dcterms:created`·`modified`** — `openpyxl` 이 저장할 때
   `utcnow()` 로 덮어쓰므로 `wb.properties` 로는 막을 수 없다. 다시 포장하면서
   **`as_of` 로 바꿔 쓴다** — 보고서의 시각은 기준 시점이어야 한다 (불변 규칙 4).

실측: 고정 전에는 1.3초 간격 두 번이 서로 달랐고, 고정 후에는 바이트 동일하다.

`openpyxl` 은 선택 의존성이다 (`.[xlsx]`). 없으면 무엇을 설치하면 되는지 말한다.
"""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from typing import List, Sequence

from .report import Report
from .safeio import UnsafeInput

#: zip 기원. 재현 가능한 빌드가 쓰는 값이다.
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_STAMP = re.compile(rb">[0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9:]{8}Z<")

_HEAD = "0F3A52"          # 머리글 배경
_WARN = "FFF4E5"          # 주의 배경

#: 엑셀 셀은 마크다운을 렌더하지 않는다. `**강조**` 가 그대로 보인다.
#: 중요한 문장을 **별표째** 보여주면 오히려 신뢰가 깎인다.
def _plain(s) -> str:
    return str(s or "").replace("**", "")


#: 긴 표준 제목 → 시트 이름. Excel 은 31자 제한에 `: \ / ? * [ ]` 를 못 쓴다.
SHEET_NAMES = {
    "KISA-ICS-2026": "KISA 점검항목",
    "NIST-SP-800-82r3-OT-OVERLAY": "NIST OT Overlay",
}


def _sheet_name(standard_id: str, title: str, n: int) -> str:
    if standard_id in SHEET_NAMES:
        return SHEET_NAMES[standard_id]
    bad = set(':\\/?*[]')
    safe = "".join(ch for ch in (standard_id or title) if ch not in bad)
    return (safe[:31] or "점검 %d" % n)


def _require():
    try:
        import openpyxl      # noqa: F401
    except ImportError:
        raise UnsafeInput(
            "엑셀 보고서를 쓰려면 openpyxl 이 필요합니다 — "
            "`python -m pip install -e \".[xlsx]\"`. 또는 `--out 보고서.html` 로 "
            "인쇄용 HTML 을 받으세요 (의존성 없음)", "openpyxl")
    import openpyxl
    return openpyxl


def _sheet(wb, title: str, headers: Sequence[str], rows: Sequence[Sequence],
           widths: Sequence[int] = ()):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet(title[:31])
    ws.append(list(headers))
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=_HEAD)
    for cell in ws[1]:
        cell.font = bold
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(list(r))
    ws.freeze_panes = "A2"
    if headers:
        ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    return ws


def _summary(wb, rep: Report):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet("요약")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 96

    def row(k, v, *, warn=False, head=False):
        ws.append([k, v])
        a, b = ws.cell(row=ws.max_row, column=1), ws.cell(row=ws.max_row, column=2)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        if head:
            a.font = Font(bold=True, color="FFFFFF")
            a.fill = b.fill = PatternFill("solid", fgColor=_HEAD)
            b.font = Font(bold=True, color="FFFFFF")
        else:
            a.font = Font(bold=True)
        if warn:
            a.fill = b.fill = PatternFill("solid", fgColor=_WARN)

    row("OT 자산 취약점·공격 경로 보고서", "", head=True)
    row("기준 시점", rep.as_of)
    row("", "이 보고서에는 '생성 시각' 이 없습니다. 기준 시점만 있고, 같은 입력이면 "
            "같은 파일이 나옵니다 — 해시를 맞춰 재현을 확인할 수 있습니다.")
    row("규칙 버전", rep.policy_version)
    row("KEV 스냅샷", rep.kev_snapshot or "미로드")
    row("권고문", "%d건" % rep.advisories)
    row("토폴로지", "%s%s" % (rep.topology_source or "없음",
                            " (합성)" if rep.topology_synthetic else
                            " (캡처에서 관측)" if rep.topology_synthetic is False else ""))
    ws.append([])

    row("자산", "%d대" % len(rep.assets), head=True)
    if rep.synthetic_assets:
        row("합성 자산", "%d대 — 실재하지 않는 장비입니다. 실제 인벤토리로 "
                      "오인하지 마세요." % rep.synthetic_assets, warn=True)
    for lvl in sorted(rep.levels):
        row("  식별 완성도 %s" % lvl, "%d대" % rep.levels[lvl])
    ws.append([])

    row("할 일", "%d건" % len(rep.actions), head=True)
    from .priority import BUCKET_ORDER, BUCKET_PHRASES
    for b in sorted(rep.buckets, key=lambda x: BUCKET_ORDER.get(x, 99)):
        row("  %s" % b, "%d건 — %s" % (rep.buckets[b], BUCKET_PHRASES.get(b, "")))
    row("통신 방식 위험", "%d건 (권고문 없이 자산만 보고 나온 것)" % len(rep.exposures))
    ws.append([])

    if rep.thin_vendors:
        row("주의 — 공개 권고문이 희박한 제조사", "", head=True)
        for c in rep.thin_vendors:
            row("  %s" % c.vendor, _plain("우리 자산 %d대 · 권고문 %d건 — %s"
                % (c.our_assets, c.advisories, c.verdict)), warn=True)
        ws.append([])

    row("이 보고서가 주장하지 않는 것", "", head=True)
    for d in rep.disclaimers:
        row("", _plain(d))
    return ws


def build_workbook(rep: Report):
    openpyxl = _require()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)                 # 기본 시트 제거
    _summary(wb, rep)

    _sheet(wb, "할 일", ["등급", "뜻", "자산", "설비명", "담당자", "구역", "제조사",
                        "모델", "권고문", "판정", "강제규칙", "KEV",
                        "확정 시 최소", "막고 있는 것", "근거"],
           [[a.bucket, a.bucket_phrase, a.asset_id, a.label, a.owner, a.zone,
             a.vendor, a.model, a.advisory_id, a.status_ko, a.fired_rules, a.kev,
             a.floor_if_confirmed, a.blocked_by, a.rationale] for a in rep.actions],
           widths=[8, 22, 20, 18, 10, 14, 20, 24, 18, 16, 12, 22, 12, 30, 50])

    _sheet(wb, "자산 대장", ["자산", "설비명", "종류", "제조사", "모델", "주문번호",
                           "펌웨어", "버전 출처", "공장", "구역", "담당자", "부서",
                           "식별 완성도", "수명주기", "주소", "합성"],
           [[a.asset_id, a.label, a.asset_type, a.vendor, a.model, a.order_number,
             a.firmware, a.firmware_method, a.factory, a.zone, a.owner,
             a.department, a.level, a.lifecycle, a.addresses, a.synthetic]
            for a in rep.assets],
           widths=[20, 18, 10, 20, 26, 20, 12, 12, 12, 14, 10, 12, 12, 12, 26, 22])

    _sheet(wb, "통신 방식 위험", ["자산", "설비명", "담당자", "종류", "코드",
                               "무슨 일인가", "그래서 뭐가 위험한가",
                               "지금 뭘 하면 되나", "등급"],
           [[e.asset_id, e.label, e.owner, e.kind, e.code, e.title, e.why,
             e.what_to_do, e.bucket] for e in rep.exposures],
           widths=[20, 18, 10, 14, 22, 40, 48, 48, 8])

    _sheet(wb, "제조사 커버리지", ["제조사", "우리 자산", "공개 권고문", "그래서"],
           [[c.vendor, c.our_assets, c.advisories, _plain(c.verdict)]
            for c in rep.coverage],
           widths=[28, 10, 12, 80])

    # 표준마다 시트를 나눈다 — 한 장에 280행을 섞으면 읽히지 않는다
    by_std: dict = {}
    for c in rep.controls:
        by_std.setdefault(c.standard, []).append(c)
    for i, (std, rows) in enumerate(sorted(by_std.items()), 1):
        name = _sheet_name(rows[0].standard_id, std, i)
        _sheet(wb, name, ["항목", "이름", "중요도", "분류", "우리 범위",
                          "근거 있음", "근거 없음", "도구가 답할 수 없음", "비고"],
               [[c.code, _plain(c.name), c.severity, c.category, c.our_scope,
                 c.present, c.absent, c.not_assessable, _plain(c.note)]
                for c in rows],
               widths=[10, 36, 8, 18, 10, 10, 10, 18, 70])
        ws = wb[name[:31]]
        ws.insert_rows(1)
        ws["A1"] = _plain("%s — 준수 여부 판정이 아닙니다. 근거를 댄 것이고, "
                          "평가는 자격을 갖춘 평가기관이 수행합니다." % std)
    return wb


def _deterministic(raw: bytes, as_of: str) -> bytes:
    """같은 입력이면 같은 바이트. zip 시각과 문서 속성 시각을 고정한다."""
    want = (">%sT00:00:00Z<" % as_of).encode("utf-8")
    src = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():          # 순서는 openpyxl 이 정한 그대로
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = _STAMP.sub(want, data)
            new = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_TIME)
            new.compress_type = zipfile.ZIP_DEFLATED
            new.external_attr = info.external_attr
            dst.writestr(new, data)
    return out.getvalue()


def write(rep: Report, path) -> Path:
    wb = build_workbook(rep)
    wb.properties.creator = "otai"
    wb.properties.lastModifiedBy = "otai"
    wb.properties.title = "OT 자산 취약점·공격 경로 보고서 (기준 시점 %s)" % rep.as_of
    buf = io.BytesIO()
    wb.save(buf)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_deterministic(buf.getvalue(), rep.as_of))
    return path

# -*- coding: utf-8 -*-
"""엑셀 자산대장을 읽는다 (ADR-046).

**현장 대장은 .xlsx 다.** CSV 만 받으면 사람이 공장별 시트를 하나씩 "CSV UTF-8"
로 다시 저장해야 하고, 그걸 매달 반복한다. 설비 800대면 그 작업이 며칠이다.

**머리글 행을 짐작하지 않는다.** 현장 대장은 1행이 병합된 제목("○○정밀 설비
자산대장"), 2행이 작성자, 3행이 빈 줄, 4행이 머리글인 경우가 흔하다. 시트마다
다르기도 하다. 그래서 이 모듈은 **후보를 점수와 함께 제시**하고, 고르는 것은
사람이다 — 추측해서 넣으면 ADR-038 의 설정값 함정과 같은 종류의 오류가 된다.
고른 결과는 매핑 프로파일에 남아 다음 달에 다시 묻지 않는다.

**xlsx 는 zip + xml 이다.** 그러므로 `safeio` 를 거친다 — 압축폭탄과 DTD 를
`openpyxl` 에 넘기기 전에 막는다. `openpyxl` 은 순수 파이썬이라 폐쇄망 반입에
플랫폼 wheel 이 필요하지 않다 (불변 규칙 9).
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from .safeio import (MAX_MEMBERS, UnsafeInput, inspect_zip, reject_dtd)

#: 시트 하나에서 머리글을 찾아볼 범위. 이보다 아래에서 시작하는 대장은 본 적이
#: 없고, 전부 훑으면 데이터 행을 머리글로 고를 위험만 커진다.
HEADER_SEARCH_ROWS = 12

#: 한 번에 읽을 시트·행 상한. CSV 와 같은 이유로 둔다 (`safeio.MAX_CSV_ROWS`).
MAX_SHEETS = 50
MAX_ROWS = 100_000


@dataclass
class HeaderCandidate:
    """머리글일 수 있는 행 하나. **확정이 아니다.**"""
    row: int                       # 1부터 센 엑셀 행 번호
    cells: Tuple[str, ...]
    recognized: Tuple[str, ...]    # 동의어 표에 걸린 머리글
    score: int                     # 걸린 개수 — 많을수록 머리글답다

    @property
    def stars(self) -> str:
        return "★" * min(3, self.score // 2) or "☆"


@dataclass
class Sheet:
    name: str
    candidates: List[HeaderCandidate] = field(default_factory=list)
    rows: int = 0

    @property
    def best(self) -> Optional[HeaderCandidate]:
        return self.candidates[0] if self.candidates else None


def _require_openpyxl():
    try:
        import openpyxl            # noqa: F401
    except ImportError:
        raise UnsafeInput(
            "엑셀을 읽으려면 openpyxl 이 필요합니다 — "
            "`python -m pip install -e \".[xlsx]\"` 로 설치하세요. "
            "또는 시트를 'CSV UTF-8' 로 저장해 `--csv` 로 주세요", "openpyxl")
    import openpyxl
    return openpyxl


def _guard(path: Path) -> None:
    """`openpyxl` 에 넘기기 전에 압축과 XML 을 본다.

    xlsx 는 zip 안에 xml 이 들어 있다. `openpyxl` 이 그 xml 을 파싱하므로,
    압축폭탄과 DTD(내부 엔티티 폭탄)는 **우리가 먼저** 막아야 한다 (ADR-038).
    """
    import zipfile
    infos = inspect_zip(path, max_members=MAX_MEMBERS)
    names = [i.filename for i in infos]
    if not any(n == "[Content_Types].xml" for n in names):
        raise UnsafeInput("xlsx 가 아닙니다 ([Content_Types].xml 이 없습니다)",
                          path.name)
    with zipfile.ZipFile(path) as zf:
        for info in infos:
            if not info.filename.endswith(".xml"):
                continue
            with zf.open(info) as fh:
                head = fh.read(1 << 16)      # 서두만 본다 — DTD 는 앞에 온다
            reject_dtd(head, member=info.filename)


def _norm(v) -> str:
    return "" if v is None else str(v).strip()


def scan(path, synonyms: Optional[Dict[str, Sequence[str]]] = None) -> List[Sheet]:
    """시트마다 머리글 **후보**를 찾는다. 아무것도 확정하지 않는다."""
    if synonyms is None:
        from .csvimport import SYNONYMS as synonyms       # 어휘의 정본은 한 곳
    known = set()
    for names in synonyms.values():
        for n in names:
            known.add(n.strip().lower().replace("_", " "))

    path = Path(path)
    _guard(path)
    openpyxl = _require_openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if len(wb.sheetnames) > MAX_SHEETS:
            raise UnsafeInput("시트 %d개 > 상한 %d" % (len(wb.sheetnames), MAX_SHEETS),
                              path.name)
        out: List[Sheet] = []
        for name in wb.sheetnames:
            ws = wb[name]
            sheet = Sheet(name=name)
            for i, raw in enumerate(
                    ws.iter_rows(max_row=HEADER_SEARCH_ROWS, values_only=True), 1):
                cells = tuple(_norm(c) for c in raw)
                filled = [c for c in cells if c]
                if len(filled) < 2:
                    continue          # 병합된 제목 한 칸은 머리글이 아니다
                hit = tuple(c for c in filled
                            if c.lower().replace("_", " ") in known)
                if hit:
                    sheet.candidates.append(
                        HeaderCandidate(row=i, cells=cells, recognized=hit,
                                        score=len(hit)))
            sheet.candidates.sort(key=lambda c: (-c.score, c.row))
            out.append(sheet)
        return out
    finally:
        wb.close()


def rows_of(path, sheet_name: str, header_row: int,
            max_rows: int = MAX_ROWS) -> Tuple[List[str], List[Tuple[int, dict]]]:
    """고른 머리글 행으로 한 시트를 읽는다.

    `(헤더, [(엑셀 행 번호, 행 dict)])` — CSV 쪽과 **같은 모양**이라 뒤 단계
    (매핑·무해화·dry-run)를 그대로 쓴다. 완전히 빈 행은 건너뛴다.
    """
    path = Path(path)
    _guard(path)
    openpyxl = _require_openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise UnsafeInput("그런 시트가 없습니다: %s (있는 것: %s)"
                              % (sheet_name, ", ".join(wb.sheetnames)), path.name)
        ws = wb[sheet_name]
        headers: List[str] = []
        rows: List[Tuple[int, dict]] = []
        for i, raw in enumerate(ws.iter_rows(values_only=True), 1):
            cells = [_norm(c) for c in raw]
            if i < header_row:
                continue
            if i == header_row:
                headers = cells
                continue
            if not any(cells):
                continue              # 현장 대장의 구분용 빈 줄
            if len(rows) >= max_rows:
                raise UnsafeInput("행 수 > 상한 %d" % max_rows, path.name)
            rows.append((i, {h: v for h, v in zip(headers, cells) if h}))
        if not headers:
            raise UnsafeInput("%d행에 머리글이 없습니다" % header_row, path.name)
        return [h for h in headers if h], rows
    finally:
        wb.close()
